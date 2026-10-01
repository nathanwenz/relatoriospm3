import streamlit as st
import pandas as pd
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import io
import os
import re
import base64
import pdfplumber
import docx
from datetime import datetime

# ReportLab para geração nativa de PDF
from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from reportlab.lib.units import inch

# =========================================================
# CONFIGURAÇÃO DA PÁGINA E TEMA (CSS)
# =========================================================
st.set_page_config(page_title="18º BPM — Gerador de Relatórios", page_icon="🛡️", layout="centered")

SENHA_CORRETA = "deusa"

def obter_caminho_brasao():
    for nome in ["brasao.png", "brasao.PNG", "Brasao.png", "BRASAO.PNG", "brasao.jpg", "brasao.jpeg"]:
        if os.path.exists(nome):
            return nome
    return None

caminho_brasao = obter_caminho_brasao()

def gerar_css_app(caminho_img):
    css_base = """
    <style>
        div[data-testid="stFileUploader"], div[data-testid="stTextInput"] {
            background-color: rgba(30, 34, 45, 0.88) !important;
            border-radius: 10px;
            padding: 10px;
            border: 1px solid #2E364A;
        }
        .stButton>button {
            background-color: #002060;
            color: #FFFFFF;
            font-weight: bold;
            border-radius: 8px;
            border: 1px solid #1E3A8A;
            width: 100%;
            padding: 10px;
        }
        .stButton>button:hover {
            background-color: #1E40AF;
            border-color: #3B82F6;
        }
        h1, h2, h3 {
            color: #F3F4F6 !important;
        }
        .developer-footer {
            position: fixed;
            bottom: 35px;
            right: 25px;
            text-align: right;
            font-size: 11px;
            font-family: sans-serif;
            color: #FFFFFF;
            background: linear-gradient(135deg, #002060, #1E3A8A);
            padding: 7px 12px;
            border-radius: 8px;
            border: 1px solid #3B82F6;
            box-shadow: 0px 4px 12px rgba(0, 0, 0, 0.4);
            z-index: 9999999 !important;
            pointer-events: none;
            line-height: 1.3;
        }
    </style>
    <div class="developer-footer">
        <span style="font-size: 9px; color: #93C5FD; text-transform: uppercase; letter-spacing: 0.5px; font-weight: bold;">Developer</span><br/>
        <strong style="font-size: 11px; color: #FFFFFF;">Nathan Wenzel</strong>
    </div>
    """
    if caminho_img and os.path.exists(caminho_img):
        try:
            with open(caminho_img, "rb") as f:
                encoded = base64.b64encode(f.read()).decode("utf-8")
            ext = caminho_img.split(".")[-1].lower()
            mime = "image/png" if "png" in ext else "image/jpeg"
            bg_css = f"""
            <style>
                .stApp {{
                    background-color: #0E1117;
                    color: #E0E6ED;
                    background-image: linear-gradient(rgba(14, 17, 23, 0.92), rgba(14, 17, 23, 0.92)), url('data:{mime};base64,{encoded}');
                    background-size: contain;
                    background-repeat: no-repeat;
                    background-position: center center;
                    background-attachment: fixed;
                }}
            </style>
            """
            return bg_css + css_base
        except Exception:
            pass

    return """
    <style>
        .stApp {
            background-color: #0E1117;
            color: #E0E6ED;
        }
    </style>
    """ + css_base

st.markdown(gerar_css_app(caminho_brasao), unsafe_allow_html=True)

# =========================================================
# CONTROLE DE SESSÃO E AUTENTICAÇÃO
# =========================================================
if "autenticado" not in st.session_state:
    st.session_state.autenticado = False

if "pagina" not in st.session_state:
    st.session_state.pagina = "inicio"

if not st.session_state.autenticado:
    if caminho_brasao:
        col1, col2, col3 = st.columns(3)
        with col2:
            st.image(caminho_brasao, width=130)
    st.title("🔒 Acesso Restrito — 18º BPM")
    st.write("Digite a senha de acesso para utilizar o Gerador de Relatórios.")

    senha_input = st.text_input("Senha de acesso:", type="password")

    if st.button("Entrar"):
        if senha_input == SENHA_CORRETA:
            st.session_state.autenticado = True
            st.success("Acesso liberado!")
            st.rerun()
        else:
            st.error("Senha incorreta! Verifique e tente novamente.")
    st.stop()


# =========================================================
# SISTEMA 1: EXTRAJORNADA VOLUNTÁRIA
# =========================================================
DIAS_SEMANA = { 0: "segunda-feira", 1: "terça-feira", 2: "quarta-feira", 3: "quinta-feira", 4: "sexta-feira", 5: "sábado", 6: "domingo" }
MESES = { 1: "janeiro", 2: "fevereiro", 3: "março", 4: "abril", 5: "maio", 6: "junho", 7: "julho", 8: "agosto", 9: "setembro", 10: "outubro", 11: "novembro", 12: "dezembro" }
MESES_REV = { "janeiro": 1, "fevereiro": 2, "março": 3, "abril": 4, "maio": 5, "junho": 6, "julho": 7, "agosto": 8, "setembro": 9, "outubro": 10, "novembro": 11, "dezembro": 12 }

def formatar_data_extenso(val):
    """Converte qualquer data para 'xx de mês de ano (dia-da-semana)' para uso no cabeçalho do relatório."""
    if val is None or pd.isna(val):
        return "23 de setembro de 2026 (quarta-feira)"
    
    if isinstance(val, (pd.Timestamp, datetime)):
        dia = val.day
        mes = MESES[val.month]
        ano = val.year
        dia_sem = DIAS_SEMANA[val.weekday()]
        return f"{dia} de {mes} de {ano} ({dia_sem})"
    
    val_str = str(val).strip()
    if not val_str or val_str.lower() in ['none', 'nan', '', 'null']:
        return "23 de setembro de 2026 (quarta-feira)"
        
    if " de " in val_str.lower() and ("feira" in val_str.lower() or "sábado" in val_str.lower() or "domingo" in val_str.lower()):
        return val_str

    s_clean = val_str.split()[0] if val_str.split() else val_str
    
    for fmt in ["%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d", "%d/%m/%y", "%Y/%m/%d"]:
        try:
            dt = datetime.strptime(s_clean, fmt)
            dia = dt.day
            mes = MESES[dt.month]
            ano = dt.year
            dia_sem = DIAS_SEMANA[dt.weekday()]
            return f"{dia} de {mes} de {ano} ({dia_sem})"
        except (ValueError, TypeError):
            pass
            
    return val_str

def formatar_data_curta(val):
    """Converte qualquer data para o formato curto numérico 'xx/yy/zzzz' para a tabela dos cards."""
    if val is None or pd.isna(val):
        return "23/09/2026"
    
    if isinstance(val, (pd.Timestamp, datetime)):
        return val.strftime("%d/%m/%Y")
        
    val_str = str(val).strip()
    if not val_str or val_str.lower() in ['none', 'nan', '', 'null']:
        return "23/09/2026"

    # Se já estiver no formato numérico
    m = re.search(r'\b(\d{1,2})[/.-](\d{1,2})[/.-](\d{2,4})\b', val_str)
    if m:
        dia, mes, ano = m.group(1).zfill(2), m.group(2).zfill(2), m.group(3)
        if len(ano) == 2:
            ano = f"20{ano}"
        return f"{dia}/{mes}/{ano}"

    # Se estiver por extenso ex: "23 de setembro de 2026 (quarta-feira)"
    m_ext = re.search(r'(\d{1,2})\s+de\s+([a-zç]+)\s+de\s+(\d{4})', val_str, re.IGNORECASE)
    if m_ext:
        dia = m_ext.group(1).zfill(2)
        mes_nome = m_ext.group(2).lower()
        ano = m_ext.group(3)
        if mes_nome in MESES_REV:
            mes_num = str(MESES_REV[mes_nome]).zfill(2)
            return f"{dia}/{mes_num}/{ano}"

    s_clean = val_str.split()[0] if val_str.split() else val_str
    for fmt in ["%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d", "%d/%m/%y", "%Y/%m/%d"]:
        try:
            dt = datetime.strptime(s_clean, fmt)
            return dt.strftime("%d/%m/%Y")
        except (ValueError, TypeError):
            pass
            
    return val_str

def formatar_horario(hora_raw):
    """Extrai estritamente o intervalo de horário, eliminando nomes de policiais ou patentes."""
    if not hora_raw or str(hora_raw).strip().lower() in ['none', 'nan', '', 'null']:
        return "Das 18:00 às 23:59"
    
    s = str(hora_raw).strip()
    
    m_range = re.search(r'(\d{1,2}(?::\d{2})?(?:\s*h)?)\s*(?:às|as|a|-|/|até|\s)\s*(\d{1,2}(?::\d{2})?(?:\s*h)?)', s, re.IGNORECASE)
    if m_range:
        h1, h2 = m_range.group(1).strip(), m_range.group(2).strip()
        if not h1.endswith('h') and ':' not in h1 and len(h1) <= 2:
            h1 = f"{h1}:00"
        if not h2.endswith('h') and ':' not in h2 and len(h2) <= 2:
            h2 = f"{h2}:00"
        return f"Das {h1} às {h2}"
        
    m_single = re.search(r'(\d{1,2}(?::\d{2})?(?:\s*h)?)', s)
    if m_single:
        h1 = m_single.group(1).strip()
        if not h1.endswith('h') and ':' not in h1 and len(h1) <= 2:
            h1 = f"{h1}:00"
        return f"Das {h1} às 23:59"
        
    return "Das 18:00 às 23:59"

def verificar_cidade_segura(file_bytes, ext, df):
    if ext == "pdf" and file_bytes is not None:
        try:
            file_bytes.seek(0)
            with pdfplumber.open(file_bytes) as pdf:
                for page in pdf.pages:
                    txt = page.extract_text() or ""
                    if "CIDADE SEGURA" in txt.upper():
                        file_bytes.seek(0)
                        return True
            file_bytes.seek(0)
        except Exception:
            pass
    if df is not None and not df.empty:
        df_str = " ".join([str(c) for c in df.columns]) + " " + " ".join([str(v) for v in df.values.flatten()])
        if "CIDADE SEGURA" in df_str.upper():
            return True
    return False

def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), hex_color)
    shd.set(qn('w:val'), 'clear')
    tcPr.append(shd)

def ler_arquivo_pdf(file_bytes):
    data = []
    file_bytes.seek(0)
    with pdfplumber.open(file_bytes) as pdf:
        for page in pdf.pages:
            tables = page.extract_tables()
            for table in tables:
                for row in table:
                    row_clean = [str(cell).replace('\n', ' ').strip() if cell is not None else '' for cell in row]
                    if any(row_clean):
                        data.append(row_clean)

        if not data:
            file_bytes.seek(0)
            for page in pdf.pages:
                tables = page.extract_tables(table_settings={
                    "vertical_strategy": "text",
                    "horizontal_strategy": "text",
                    "snap_tolerance": 5,
                })
                for table in tables:
                    for row in table:
                        row_clean = [str(cell).replace('\n', ' ').strip() if cell is not None else '' for cell in row]
                        if any(row_clean):
                            data.append(row_clean)

    if not data:
        return pd.DataFrame()

    header_idx = -1
    for idx, row in enumerate(data):
        row_str = " ".join(row).upper()
        if any(term in row_str for term in ["VOLCHER", "VOUCHER", "CIDADE", "MUNICÍPIO", "MUNICIPIO", "DATA", "HORA", "HORÁRIO", "N°", "Nº"]):
            header_idx = idx
            break

    if header_idx == -1:
        headers = [f"COL_{i}" for i in range(len(data))] if data else []
        rows = data
    else:
        headers = [str(h).strip() if str(h).strip() else f"COL_{i}" for i, h in enumerate(data[header_idx])]
        rows = data[header_idx + 1:]

    rows_filtradas = []
    for row in rows:
        row_str = " ".join(row).upper()
        if "VOLCHER" in row_str and "CIDADE" in row_str:
            continue
        if "VOUCHER" in row_str and "CIDADE" in row_str:
            continue
        if "N°" in row_str and "CIDADE" in row_str:
            continue
        rows_filtradas.append(row)

    return pd.DataFrame(rows_filtradas, columns=headers)

def processar_dataframe(df):
    if df is None or df.empty:
        return df, None, None, None, None

    cols = {str(c).upper().strip(): c for c in df.columns}
    
    def normalizar(txt):
        txt = str(txt).upper()
        return re.sub(r'[ÁÀÂÃ]', 'A', re.sub(r'[ÉÈÊ]', 'E', re.sub(r'[ÍÌÎ]', 'I', re.sub(r'[ÓÒÔÕ]', 'O', re.sub(r'[ÚÙÛ]', 'U', txt)))))

    termos_policial = ["NOME", "POLICIAL", "EFETIVO", "GRADUACAO", "POSTO", "GRAD", "PM", "MILITAR", "INTEGRANTE", "RG", "CPF", "MATRICULA", "CONTATO", "TELEFONE", "RESPONSAVEL", "SERVIDOR"]

    def e_coluna_policial(col_nome):
        col_norm = normalizar(col_nome)
        return any(tp in col_norm for tp in termos_policial)

    def encontrar_coluna_por_nome(termos, df_cols, ignorar_policial=False):
        for termo in termos:
            termo_norm = normalizar(termo)
            for col_upper, col_orig in df_cols.items():
                col_norm = normalizar(col_upper)
                if ignorar_policial and e_coluna_policial(col_orig):
                    continue
                if termo_norm in col_norm:
                    return col_orig
        return None

    col_volcher = encontrar_coluna_por_nome(["VOLCHER", "VOUCHER", "VOLCHE", "VOUCHE", "N°", "Nº", "NUMERO", "NRO", "CODIGO", "CARTAO"], cols)
    col_cidade = encontrar_coluna_por_nome(["CIDADE", "MUNICÍPIO", "MUNICIPIO", "LOCAL", "OPM", "MUNIC", "LOTAÇÃO", "POSTO"], cols)
    col_data = encontrar_coluna_por_nome(["DATA", "DIA", "PERIODO"], cols, ignorar_policial=True)
    col_hora = encontrar_coluna_por_nome(["HORA", "HORARIO", "TURNO", "PERIODO", "INICIO", "ESCALA", "SERVICO", "TEMPO"], cols, ignorar_policial=True)

    for col in df.columns:
        if e_coluna_policial(col):
            continue
            
        sample_vals = [str(v) for v in df[col].dropna().head(10).tolist()]
        sample_text = " ".join(sample_vals)
        
        if col_hora is None and col not in [col_volcher, col_cidade, col_data]:
            if re.search(r'\b\d{1,2}(?::\d{2}|h)\b', sample_text, re.IGNORECASE) or "AS" in sample_text.upper() or "ÀS" in sample_text.upper():
                col_hora = col

        if col_volcher is None and col not in [col_cidade, col_data, col_hora]:
            if any(re.search(r'^\d{1,5}$', v.strip()) for v in sample_vals):
                col_volcher = col

        if col_data is None and col not in [col_volcher, col_cidade, col_hora]:
            if re.search(r'\b\d{1,2}/\d{1,2}/\d{2,4}\b', sample_text) or re.search(r'\b\d{4}-\d{2}-\d{2}\b', sample_text):
                col_data = col

    return df, col_volcher, col_cidade, col_data, col_hora

def gerar_relatorio_word(df_escala, data_extenso="23 de setembro de 2026 (quarta-feira)", e_cidade_segura=False):
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(0.6)
        section.bottom_margin = Inches(0.6)
        section.left_margin = Inches(0.7)
        section.right_margin = Inches(0.7)

    p_hdr = doc.add_paragraph()
    p_hdr.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_hdr.paragraph_format.space_after = Pt(2)
    p_hdr.paragraph_format.line_spacing = 1.15

    r_estado = p_hdr.add_run("ESTADO DO PARANÁ\nPOLÍCIA MILITAR\n2º COMANDO REGIONAL DE POLÍCIA MILITAR\n18° BATALHÃO DE POLÍCIA MILITAR\n")
    r_estado.bold = True
    r_estado.font.size = Pt(11)
    r_estado.font.name = "Arial"

    r_prog = p_hdr.add_run("PROGRAMAÇÃO EXTRAJORNADA VOLUNTÁRIA\n")
    r_prog.bold = True
    r_prog.font.size = Pt(12)
    r_prog.font.name = "Arial"
    r_prog.font.color.rgb = RGBColor(0, 32, 96)

    if e_cidade_segura:
        r_cs = p_hdr.add_run("CIDADE SEGURA\n")
        r_cs.bold = True
        r_cs.font.size = Pt(12)
        r_cs.font.name = "Arial"
        r_cs.font.color.rgb = RGBColor(0, 32, 96)

    # Data no cabeçalho principal do documento (por extenso)
    data_cabecalho_formatada = formatar_data_extenso(data_extenso)

    p_data = doc.add_paragraph()
    p_data.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_data.paragraph_format.space_after = Pt(6)
    r_dt = p_data.add_run(data_cabecalho_formatada)
    r_dt.bold = True
    r_dt.font.size = Pt(11)
    r_dt.font.name = "Arial"

    p_obs = doc.add_paragraph()
    p_obs.paragraph_format.space_after = Pt(12)
    r_obs = p_obs.add_run("* As equipes ficarão à disposição do CPU, Adjunto e COPOM para atendimentos de ocorrências. Na ausência de ocorrências, deverão seguir o cartão de programa.")
    r_obs.italic = True
    r_obs.font.size = Pt(9.5)
    r_obs.font.name = "Arial"

    df, col_v, col_c, col_d, col_h = processar_dataframe(df_escala)

    for c in [col_v, col_c, col_d, col_h]:
        if c and c in df.columns:
            df[c] = df[c].astype(str).str.strip()

    group_cols = [c for c in [col_v, col_c, col_d, col_h] if c is not None and c in df.columns]

    if group_cols:
        grupos_iterator = df.groupby(group_cols, sort=False)
    else:
        grupos_iterator = enumerate(df.iterrows())

    for group_idx, item in enumerate(grupos_iterator):
        if group_cols:
            chaves, grupo = item
            primeiro = grupo.iloc
        else:
            _, row = item
            primeiro = row

        volcher_raw = str(primeiro.get(col_v, "")).replace('.0', '').replace('None', '').replace('nan', '').strip() if col_v else ""
        volcher_val = volcher_raw if volcher_raw else str(group_idx + 1)

        cidade_val = str(primeiro.get(col_c, "")).replace('None', '').replace('nan', '').strip() if col_c else ""
        
        # Data dentro da tabela do card em formato curto ex: 23/09/2026
        data_raw = str(primeiro.get(col_d, "")).strip() if col_d else ""
        data_str = formatar_data_curta(data_raw) if data_raw else formatar_data_curta(data_cabecalho_formatada)

        hora_raw = str(primeiro.get(col_h, "")).strip() if col_h else ""
        hora_str = formatar_horario(hora_raw)

        table = doc.add_table(rows=5, cols=2)
        table.style = 'Table Grid'
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False

        for row in table.rows:
            row.cells.width = Inches(2.0)
            row.cells.width = Inches(4.5)

        campos = [
            ("CIDADE/VOLCHER", f"{cidade_val} - VOLCHER - {volcher_val}"),
            ("DATA", data_str),
            ("LOCAL", cidade_val),
            ("HORÁRIO", hora_str)
        ]

        for i, (label, val) in enumerate(campos):
            r = table.rows[i]

            c0 = r.cells
            p0 = c0.paragraphs
            p0.paragraph_format.space_after = Pt(2)
            p0.paragraph_format.space_before = Pt(2)
            r0 = p0.add_run(label)
            r0.bold = True
            r0.font.name = "Arial"
            r0.font.size = Pt(10)
            set_cell_background(c0, "D9E1F2")

            c1 = r.cells
            p1 = c1.paragraphs
            p1.paragraph_format.space_after = Pt(2)
            p1.paragraph_format.space_before = Pt(2)
            r1 = p1.add_run(val)
            r1.bold = True if label in ["CIDADE/VOLCHER", "LOCAL"] else False
            r1.font.name = "Arial"
            r1.font.size = Pt(10)
            set_cell_background(c1, "FFFFFF")

        r4 = table.rows
        c0 = r4.cells
        c1 = r4.cells
        c0.merge(c1)
        p_obs_tbl = c0.paragraphs
        p_obs_tbl.paragraph_format.space_after = Pt(3)
        p_obs_tbl.paragraph_format.space_before = Pt(3)
        p_obs_tbl.paragraph_format.line_spacing = 1.15

        r_l1 = p_obs_tbl.add_run("A equipe ficará a Disposição do COPOM e CPU ou Adjunto.\n")
        r_l1.font.name = "Arial"
        r_l1.font.size = Pt(9.5)

        r_l2 = p_obs_tbl.add_run("SISGCOP 61076\n")
        r_l2.bold = True
        r_l2.font.name = "Arial"
        r_l2.font.size = Pt(9.5)

        r_l3 = p_obs_tbl.add_run("    • Equipe deverá fazer contato com o Adjunto ao assumir serviço.\n")
        r_l3.font.name = "Arial"
        r_l3.font.size = Pt(9.5)

        r_l4 = p_obs_tbl.add_run("    • A equipe além realizar o atendimento de ocorrências, deverá realizar o Patrulhamento Ostensivo e Preventivo na área designada para atuar.")
        r_l4.font.name = "Arial"
        r_l4.font.size = Pt(9.5)

        set_cell_background(c0, "FAFAFA")

        doc.add_paragraph().paragraph_format.space_after = Pt(4)

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer

def render_extrajornada():
    col_back, _ = st.columns(2)
    with col_back:
        if st.button("⬅️ Voltar ao Início", key="btn_voltar_extrajornada"):
            st.session_state.pagina = "inicio"
            st.rerun()

    if caminho_brasao:
        col1, col2, col3 = st.columns(3)
        with col2:
            st.image(caminho_brasao, width=120)

    st.title("🛡️ Gerador de Relatório Extrajornada")
    st.caption("18º Batalhão de Polícia Militar — PMPR")
    st.write("Envie a tabela da escala em **.pdf**, **.xlsx** ou **.csv**.")

    arquivo = st.file_uploader("Envie o arquivo da escala", type=["xlsx", "csv", "pdf"])
    if arquivo:
        ext = arquivo.name.split(".")[-1].lower()
        if ext == "xlsx":
            df = pd.read_excel(arquivo)
        elif ext == "csv":
            df = pd.read_csv(arquivo)
        elif ext == "pdf":
            df = ler_arquivo_pdf(arquivo)
        else:
            df = pd.DataFrame()

        if not df.empty:
            st.success(f"Arquivo carregado com sucesso! {len(df)} registros encontrados.")
            
            e_cidade_segura = verificar_cidade_segura(arquivo, ext, df)
            if e_cidade_segura:
                st.info("ℹ️ Operação 'CIDADE SEGURA' identificada no arquivo. O cabeçalho do relatório incluirá este destaque.")

            _, _, _, col_d, _ = processar_dataframe(df)
            data_sugerida_tela = "23 de setembro de 2026 (quarta-feira)"
            if col_d and col_d in df.columns and not df[col_d].dropna().empty:
                primeira_data_val = df[col_d].dropna().iloc
                data_sugerida_tela = formatar_data_extenso(primeira_data_val)

            data_cabecalho = st.text_input("Data para o cabeçalho do relatório", value=data_sugerida_tela)
            
            if st.button("Gerar Documento Word"):
                docx_bytes = gerar_relatorio_word(df, data_cabecalho, e_cidade_segura)
                st.download_button("📥 Baixar Relatório Preenchido (.docx)", docx_bytes, "RELATORIO_EXTRAJORNADA.docx")
        else:
            st.error("Não foi possível extrair dados da tabela. Verifique o arquivo enviado.")


# =========================================================
# SISTEMA 2: ORDEM DE SERVIÇO (OO -> OS)
# =========================================================
def data_atual_extenso():
    now = datetime.now()
    return f"{now.day} de {MESES[now.month]} de {now.year}"

def set_cell_bg(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def limpar_assinaturas_e_ruidos(texto):
    if not texto:
        return ""
    padroes_lixo = [
        r'.*Assinado\s+eletronicamente.*',
        r'Documento\s+assinado\s+digitalmente.*',
        r'Inserido\s+ao\s+protocolo.*',
        r'conforme\s+MP\s+n[º°\.]?\s*2\.?200-2/2001.*',
        r'Validação:.*',
        r'https?://[^\s]*eprotocolo[^\s]*',
        r'www\.[^\s]*eprotocolo[^\s]*',
        r'Código\s+de\s+autenticidade:.*',
        r'Assinado\s+eletronicamente\s+por:.*',
        r'Assinatura\s+Qualificada\s+efetuada\s+por:.*',
        r'Assinatura\s+digital.*',
        r'Chave\s+de\s+Autenticação:.*',
        r'Para\s+verificar\s+a\s+autenticidade.*',
        r'e-Protocolo\s*\d+.*',
        r'SHA-?256:.*',
        r'Hash\s*:.*',
        r'Demais\s+assinaturas\s+na\s+folha.*',
        r'A\s+autenticidade\s+deste.*',
        r'\d{1,2}\.\d{3}\.\d{3}-\d.*por:.*',
        r'.*por:\s*\d*º?\s*(Sgt|Ten|Cel|Cap|Maj|Subten|Cb|Sd|Aux|PM).*em:\s*\d{2}/\d{2}/\d{4}.*',
        r'.*em:\s*\d{2}/\d{2}/\d{4}\s*\d{2}:\d{2}.*',
        r'Folha\s+\d+[a-z]?\s*de\s*\d+.*',
        r'Página\s+\d+\s+de\s+\d+.*',
        r'Aux\.\s*P/\d+.*em:\s*\d{2}/\d{2}/\d{4}.*',
        r'.*documento\s+pode\s+ser\s+verificad[oa].*',
        r'.*site\s+do\s+eprotocolo.*',
        r'^\s*\d{1,3}[a-z]?\s*$',
        r'^\s*-\s*\d{1,3}\s*-\s*$',
        r'^\s*\d{2}:\d{2}\b.*',
    ]

    linhas = texto.split('\n')
    linhas_limpas = []

    for linha in linhas:
        l_trim = linha.strip()
        if not l_trim:
            continue
            
        descartar = False
        for p in padroes_lixo:
            if re.search(p, l_trim, re.IGNORECASE):
                descartar = True
                break
                
        if not descartar:
            if '|' not in l_trim:
                l_trim = re.sub(r'\s+\d{1,2}$', '', l_trim)
            linhas_limpas.append(l_trim)
            
    res = '\n'.join(linhas_limpas)
    res = re.sub(r'\n{3,}', '\n\n', res)
    return res.strip()

def formatar_quebras_de_secao(texto):
    if not texto:
        return ""
    linhas = texto.split('\n')
    novas_linhas = []
    for l in linhas:
        if '|' in l:
            novas_linhas.append(l)
        else:
            l = re.sub(r'([^\n])\s*(\d+[ªº]\s*FASE|FASE\s+\d+|FASE\s+[I|V|X]+)', r'\1\n\2', l, flags=re.IGNORECASE)
            l = re.sub(r'([^\n])\s+([a-z0-9]{1,3}[.)])\s+', r'\1\n\2 ', l, flags=re.IGNORECASE)
            l = re.sub(r'([^\n])\s+(\d+.\d+(?:.\d+)?)\s+', r'\1\n\2 ', l)
            novas_linhas.append(l)
    return '\n'.join(novas_linhas)

def safe_crop_text(page, y0, y1):
    if y1 <= y0 + 1:
        return ""
    y0 = max(0, min(y0, page.height - 1))
    y1 = max(y0 + 1, min(y1, page.height))
    if y1 <= y0:
        return ""
    try:
        cropped = page.crop((0, y0, page.width, y1))
        return cropped.extract_text() or ""
    except Exception:
        return ""

def is_real_data_table(extracted_tbl):
    if not extracted_tbl or len(extracted_tbl) == 0:
        return False
    max_cols = 0
    for row in extracted_tbl:
        if row:
            non_empty = [c for c in row if c is not None and str(c).strip() != ""]
            if len(non_empty) > max_cols:
                max_cols = len(non_empty)
    return max_cols >= 2

def extrair_texto_arquivo(uploaded_file):
    ext = uploaded_file.name.split(".")[-1].lower()
    text = ""
    if ext == "pdf":
        uploaded_file.seek(0)
        with pdfplumber.open(uploaded_file) as pdf:
            for page in pdf.pages:
                tables = page.find_tables()
                if not tables:
                    page_text = page.extract_text() or ""
                    if page_text.strip():
                        text += page_text + "\n\n"
                else:
                    table_bboxes = [t.bbox for t in tables]
                    
                    def not_in_table(obj):
                        obj_x0 = obj.get('x0', 0)
                        obj_top = obj.get('top', 0)
                        obj_x1 = obj.get('x1', 0)
                        obj_bottom = obj.get('bottom', 0)
                        for (tx0, ttop, tx1, tbottom) in table_bboxes:
                            if not (obj_x1 <= tx0 or obj_x0 >= tx1 or obj_bottom <= ttop or obj_top >= tbottom):
                                return False
                        return True

                    filtered_page = page.filter(not_in_table)
                    tables_sorted = sorted(tables, key=lambda t: t.bbox)
                    
                    page_content = []
                    last_top = 0
                    page_height = page.height
                    
                    for t in tables_sorted:
                        tx0, ttop, tx1, tbottom = t.bbox
                        if ttop > last_top + 2:
                            t_above = safe_crop_text(filtered_page, last_top, ttop)
                            if t_above and t_above.strip():
                                page_content.append(t_above.strip())
                        
                        extracted_tbl = t.extract()
                        if extracted_tbl:
                            if is_real_data_table(extracted_tbl):
                                table_lines = []
                                for row in extracted_tbl:
                                    if row and any(c is not None and str(c).strip() != "" for c in row):
                                        clean_row = [str(c).replace('\n', ' ').strip() if c is not None else "" for c in row]
                                        table_lines.append(" | ".join(clean_row))
                                if table_lines:
                                    page_content.append("\n".join(table_lines))
                            else:
                                plain_lines = []
                                for row in extracted_tbl:
                                    for cell in row:
                                        if cell and str(cell).strip():
                                            plain_lines.append(str(cell).strip())
                                if plain_lines:
                                    page_content.append("\n".join(plain_lines))
                        
                        last_top = max(last_top, tbottom)
                    
                    if last_top < page_height - 2:
                        t_below = safe_crop_text(filtered_page, last_top, page_height)
                        if t_below and t_below.strip():
                            page_content.append(t_below.strip())
                            
                    text += "\n\n".join(page_content) + "\n\n"
        uploaded_file.seek(0)
    elif ext in ["docx", "doc"]:
        uploaded_file.seek(0)
        doc = docx.Document(uploaded_file)
        items = []
        for element in doc.element.body:
            if element.tag.endswith('p'):
                p = docx.text.paragraph.Paragraph(element, doc)
                if p.text.strip():
                    items.append(p.text.strip())
            elif element.tag.endswith('tbl'):
                tbl = docx.table.Table(element, doc)
                table_matrix = []
                for row in tbl.rows:
                    row_cells = [c.text.replace('\n', ' ').strip() for c in row.cells]
                    if any(cell for cell in row_cells):
                        table_matrix.append(row_cells)
                
                if table_matrix:
                    if is_real_data_table(table_matrix):
                        table_lines = [" | ".join(r) for r in table_matrix]
                        items.append("\n".join(table_lines))
                    else:
                        plain_lines = [" ".join([c for c in r if c]) for r in table_matrix]
                        items.append("\n".join(plain_lines))
        text = "\n\n".join(items)
        uploaded_file.seek(0)
        
    text_limpo = limpar_assinaturas_e_ruidos(text)
    return text_limpo

def extrair_secao_flexivel(texto, padrao_inicio, padraos_fim):
    regex_inicio = rf'^[ \t]*[\d.]*\s*{padrao_inicio}[^\n]*\n'
    match_inicio = re.search(regex_inicio, texto, re.IGNORECASE | re.MULTILINE)

    if not match_inicio:
        regex_fb = rf'{padrao_inicio}'
        match_fb = re.search(regex_fb, texto, re.IGNORECASE)
        if not match_fb:
            return ""
        pos_inicio = match_fb.end()
        nl = texto.find('\n', pos_inicio)
        pos_conteudo = nl + 1 if nl != -1 else pos_inicio
    else:
        pos_conteudo = match_inicio.end()
        
    conteudo_restante = texto[pos_conteudo:]

    regex_fim = rf'^[ \t]*[\d\.]*\s*(?:{"|".join(padraos_fim)})[^\n]*'
    match_fim = re.search(regex_fim, conteudo_restante, re.IGNORECASE | re.MULTILINE)

    if match_fim:
        conteudo_secao = conteudo_restante[:match_fim.start()]
    else:
        conteudo_secao = conteudo_restante
        
    return conteudo_secao.strip()

def normalizar_subnumeracao_secao(conteudo, sec_num):
    if not conteudo:
        return ""

    linhas = conteudo.strip().split('\n')
    linhas_limpas = []

    padrao_titulo_principal = re.compile(
        r'^\s*(\d*[\.\)]?\s*)?(E\s+LOGÍSTICA|FINALIDADE|INFORMAÇÕES\s+GERAIS|SITUAÇÃO|MISSÃO|EXECUÇÃO|ADMINISTRAÇÃO|LOGÍSTICA|RELATÓRIOS|PRESCRIÇÕES\s+DIVERSAS|REFERÊNCIAS)\b.*$',
        re.IGNORECASE
    )

    padrao_numero_isolado = re.compile(r'^\s*\d+\s*[\.\)]?\s*$')

    sub_counter = 1
    seen_subnums = set()

    for linha in linhas:
        l_str = linha.strip()
        if not l_str:
            continue

        if padrao_titulo_principal.match(l_str) and len(l_str) < 70 and not '|' in l_str:
            continue

        if padrao_numero_isolado.match(l_str):
            continue

        match_sub = re.match(r'^(\d+)\.(\d+)(\.?)\s*(.*)', l_str)
        if match_sub:
            prefix_main, prefix_sub, prefix_dot, rest = match_sub.groups()

            curr_sub_num = int(prefix_sub)
            sub_key = f"{sec_num}.{curr_sub_num}"

            if int(prefix_main) != sec_num or sub_key in seen_subnums:
                new_sub_str = f"{sec_num}.{sub_counter}. {rest}".strip()
                seen_subnums.add(f"{sec_num}.{sub_counter}")
                sub_counter += 1
            else:
                new_sub_str = f"{sec_num}.{curr_sub_num}. {rest}".strip()
                seen_subnums.add(sub_key)
                sub_counter = max(sub_counter, curr_sub_num + 1)

            linhas_limpas.append(new_sub_str)
        else:
            linhas_limpas.append(l_str)

    return "\n".join(linhas_limpas).strip()

def parsear_ordem_operacao(texto):
    dados = {}

    match_num = re.search(r'ORDEM DE OPERAÇÃO\s*(?:Nº|N°|Nº\.|N°\.|N°\s*|Nº\s*)?(\d+/\d+)', texto, re.IGNORECASE)
    dados['num_oo'] = match_num.group(1) if match_num else "000/2026"

    match_nome = re.search(r'“([^”]+)”|"([^"]+)"', texto)
    if match_nome:
        dados['nome_op'] = match_nome.group(1) or match_nome.group(2)
    else:
        match_nome2 = re.search(r'OPERAÇÃO\s+([A-Z0-9\s–\-]{4,})', texto)
        dados['nome_op'] = match_nome2.group(0).strip() if match_nome2 else "OPERAÇÃO POLICIAL"

    dados['nome_op'] = dados['nome_op'].replace('\n', ' ').strip().upper()

    f = extrair_secao_flexivel(texto, r'FINALIDADE', [r'SITUAÇÃO', r'INFORMAÇÕES\s+GERAIS', r'MISSÃO'])
    f = normalizar_subnumeracao_secao(f, 1)
    dados['finalidade'] = f if f else "Realizar ações de policiamento ostensivo preventivo e preservação da ordem pública."

    ig_raw = extrair_secao_flexivel(texto, r'SITUAÇÃO|INFORMAÇÕES\s+GERAIS', [r'MISSÃO', r'EXECUÇÃO'])
    lines_ig = []
    for line in ig_raw.split('\n'):
        l_str = line.strip()
        if not l_str:
            continue
        if re.match(r'^\s*(\d*[\.\)]?\s*)?(SITUAÇÃO|INFORMAÇÕES\s+GERAIS)\s*$', l_str, re.IGNORECASE):
            continue
        if re.match(r'^\s*\d+\.\d+\.?\s*(INFORMAÇÕES\s+GERAIS|SITUAÇÃO)\s*$', l_str, re.IGNORECASE):
            continue
        lines_ig.append(l_str)
        
    ig_body = "\n".join(lines_ig).strip()
    ig_final = "2.1. SITUAÇÃO\n" + ig_body
    ig_final = normalizar_subnumeracao_secao(ig_final, 2)
    dados['informacoes_gerais'] = ig_final

    m = extrair_secao_flexivel(texto, r'MISSÃO', [r'EXECUÇÃO', r'ADMINISTRAÇÃO', r'LOGÍSTICA'])
    if not m or len(m) < 10:
        m = f"O 18º BPM executará o policiamento ostensivo e a preservação da ordem pública na sua circunscrição territorial no âmbito da “{dados['nome_op']}”, visando a prevenção de crimes e a garantia da segurança pública."
    else:
        m = normalizar_subnumeracao_secao(m, 3)
    dados['missao'] = m

    e = extrair_secao_flexivel(texto, r'EXECUÇÃO', [r'ADMINISTRAÇÃO', r'LOGÍSTICA', r'RELATÓRIOS'])
    e = normalizar_subnumeracao_secao(e, 4)
    dados['execucao'] = e if e else "Atuação integrada das equipes operacionais do 18º BPM em conformidade com o planejamento."

    l = extrair_secao_flexivel(texto, r'ADMINISTRAÇÃO|LOGÍSTICA', [r'RELATÓRIOS', r'PRESCRIÇÕES'])
    l = normalizar_subnumeracao_secao(l, 5)
    dados['logistica'] = l if l else "Uniforme: Orgânico da OPM (4º RUPM).\nArmamento e equipamento: Orgânico compatível com o serviço.\nTransporte: Viaturas operacionais do 18º BPM."

    r = extrair_secao_flexivel(texto, r'RELATÓRIOS', [r'PRESCRIÇÕES', r'REFERÊNCIAS'])
    match_sisgcop = re.search(r'(\d{5,6})\s*[\-–]?\s*[\"“]?OPERAÇÃO', texto, re.IGNORECASE)
    if not match_sisgcop:
        match_sisgcop = re.search(r'SISGCOP[^\d]*(\d{5,6})', texto, re.IGNORECASE)
    num_sisgcop = match_sisgcop.group(1) if match_sisgcop else ""

    if r:
        r = normalizar_subnumeracao_secao(r, 6)
        dados['relatorios'] = r
    else:
        dados['relatorios'] = f"Os resultados obtidos deverão ser lançados no SISGCOP{' sob o código ' + num_sisgcop if num_sisgcop else ''} até o término da operação. Confecção dos Boletins de Ocorrência (BOU) no SADE."

    p = extrair_secao_flexivel(texto, r'PRESCRIÇÕES\s+DIVERSAS|PRESCRIÇÕES', [r'REFERÊNCIAS', r'DISTRIBUIÇÃO'])
    if p:
        p = normalizar_subnumeracao_secao(p, 7)
        dados['prescricoes'] = p
    else:
        dados['prescricoes'] = "Os policiais militares deverão atuar com bom senso, urbanidade, legalidade e estrito cumprimento do dever legal. Preleção obrigatória antes do início do serviço."

    ref = extrair_secao_flexivel(texto, r'REFERÊNCIAS', [r'DISTRIBUIÇÃO', r'$'])
    if ref and len(ref.strip()) > 10:
        dados['referencias'] = ref
    else:
        dados['referencias'] = f"a) Constituição da República Federativa do Brasil de 1988;\nb) Constituição do Estado do Paraná de 1989;\nc) Lei n.º 22.354/2025 – Lei de Organização Básica da PMPR;\nd) Ordem de Operação nº {dados['num_oo']} – 2º CRPM ({dados['nome_op']});\ne) Determinação do Comandante do 18º BPM."

    return dados

def eh_titulo_subsecao(linha):
    l = linha.strip()
    if not l:
        return False
    if l.endswith(':') and len(l) < 90:
        return True
    if re.match(r'^\d+(.\d+)+[.)-]?\s+', l):
        return True
    if re.match(r'^(FASE|ETAPA|GRUPO|ROTA|ZONA)\s+', l, re.IGNORECASE):
        return True
    if re.match(r'^[I|V|X]+[.)-]\s+', l, re.IGNORECASE):
        return True
    if l.isupper() and len(l) < 90 and not '|' in l and not l.startswith('PMPR'):
        return True
    return False

def renderizar_conteudo_docx(doc, conteudo, forcar_texto=False):
    conteudo_formatado = formatar_quebras_de_secao(conteudo)
    if not conteudo_formatado or not conteudo_formatado.strip():
        return
    if forcar_texto:
        conteudo_formatado = conteudo_formatado.replace('|', ' ')

    linhas = conteudo_formatado.strip().split('\n')
    i = 0
    while i < len(linhas):
        linha = linhas[i].strip()
        if not linha:
            i += 1
            continue
            
        if not forcar_texto and '|' in linha:
            tabela_linhas = []
            while i < len(linhas) and '|' in linhas[i]:
                cels = [c.strip() for c in linhas[i].split('|')]
                if len(cels) > 1 and cels == "":
                    cels = cels[1:]
                if len(cels) > 1 and cels[-1] == "":
                    cels = cels[:-1]
                if any(c for c in cels):
                    tabela_linhas.append(cels)
                i += 1
                
            if tabela_linhas:
                max_cols = max(len(r) for r in tabela_linhas)
                if max_cols < 2:
                    for r_data in tabela_linhas:
                        txt_line = " ".join([c for c in r_data if c])
                        if txt_line.strip():
                            p = doc.add_paragraph()
                            p.paragraph_format.line_spacing = 1.2
                            p.paragraph_format.space_before = Pt(0)
                            p.paragraph_format.space_after = Pt(6)
                            r = p.add_run(txt_line.strip())
                            r.font.name = "Arial"
                            r.font.size = Pt(10)
                else:
                    for r_data in tabela_linhas:
                        while len(r_data) < max_cols:
                            r_data.append("")

                    tbl = doc.add_table(rows=len(tabela_linhas), cols=max_cols)
                    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                    tbl.style = 'Table Grid'
                    
                    for r_idx, row_data in enumerate(tabela_linhas):
                        row_cells = tbl.rows[r_idx].cells
                        for c_idx, cell_value in enumerate(row_data):
                            if c_idx < len(row_cells):
                                cell = row_cells[c_idx]
                                p = cell.paragraphs if cell.paragraphs else cell.add_paragraph()
                                p.paragraph_format.space_before = Pt(3)
                                p.paragraph_format.space_after = Pt(3)
                                p.paragraph_format.line_spacing = 1.15
                                
                                r = p.add_run(cell_value)
                                r.font.name = "Arial"
                                
                                if r_idx == 0:
                                    set_cell_bg(cell, "002060")
                                    r.bold = True
                                    r.font.size = Pt(9.5)
                                    r.font.color.rgb = RGBColor(255, 255, 255)
                                else:
                                    if r_idx % 2 == 1:
                                        set_cell_bg(cell, "F8FAFC")
                                    r.font.size = Pt(9.0)
                                    r.font.color.rgb = RGBColor(30, 41, 59)
                    p_sp = doc.add_paragraph()
                    p_sp.paragraph_format.space_after = Pt(4)
            continue
            
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.2

        if eh_titulo_subsecao(linha):
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            r = p.add_run(linha)
            r.bold = True
            r.font.name = "Arial"
            r.font.size = Pt(10.5)
        else:
            m_item = re.match(r'^([a-z0-9]{1,3}[\.\)\-]|[A-Z][\.\)\-])\s+(.*)', linha, re.IGNORECASE)
            if m_item:
                prefix, rest = m_item.groups()
                p.paragraph_format.space_before = Pt(3)
                p.paragraph_format.space_after = Pt(4)
                
                r_pre = p.add_run(prefix + " ")
                r_pre.bold = True
                r_pre.font.name = "Arial"
                r_pre.font.size = Pt(10)
                
                r_rest = p.add_run(rest)
                r_rest.font.name = "Arial"
                r_rest.font.size = Pt(10)
            else:
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(6)
                r = p.add_run(linha)
                r.font.name = "Arial"
                r.font.size = Pt(10)
        
        i += 1

def gerar_ordem_servico_docx(fields):
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    table_hdr = doc.add_table(rows=1, cols=2)
    table_hdr.autofit = False
    table_hdr.alignment = WD_TABLE_ALIGNMENT.CENTER

    row0 = table_hdr.rows
    cell_left = row0.cells
    cell_right = row0.cells

    cell_left.width = Inches(3.5)
    cell_right.width = Inches(3.0)

    p_left = cell_left.paragraphs
    p_left.paragraph_format.space_after = Pt(2)
    p_left.paragraph_format.line_spacing = 1.2
    r_l = p_left.add_run("PMPR\n2º CRPM/18º BPM\nP/3")
    r_l.bold = True
    r_l.font.name = "Arial"
    r_l.font.size = Pt(10)

    p_right = cell_right.paragraphs
    p_right.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_right.paragraph_format.space_after = Pt(2)
    p_right.paragraph_format.line_spacing = 1.2
    r_r = p_right.add_run(f"Cornélio Procópio, PR.\nEm {fields.get('data_expedicao', data_atual_extenso())}\nORDEM DE SERVIÇO Nº {fields.get('num_os', '077')}")
    r_r.bold = True
    r_r.font.name = "Arial"
    r_r.font.size = Pt(10)

    p_div = doc.add_paragraph()
    p_div.paragraph_format.space_after = Pt(10)
    p_div.paragraph_format.space_before = Pt(6)
    r_div = p_div.add_run("___________________________________________________________________")
    r_div.bold = True
    r_div.font.size = Pt(9)

    p_titulo = doc.add_paragraph()
    p_titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_titulo.paragraph_format.space_before = Pt(6)
    p_titulo.paragraph_format.space_after = Pt(16)
    r_tit = p_titulo.add_run(f"“{fields.get('nome_operacao', '').strip().upper()}”")
    r_tit.bold = True
    r_tit.font.name = "Arial"
    r_tit.font.size = Pt(12)
    r_tit.font.color.rgb = RGBColor(0, 32, 96)

    secoes = [
        ("1. FINALIDADE", fields.get('finalidade', '')),
        ("2. INFORMAÇÕES GERAIS", fields.get('informacoes_gerais', '')),
        ("3. MISSÃO", fields.get('missao', '')),
        ("4. EXECUÇÃO", fields.get('execucao', '')),
        ("5. ADMINISTRAÇÃO E LOGÍSTICA", fields.get('logistica', '')),
        ("6. RELATÓRIOS E SISGCOP", fields.get('relatorios', '')),
        ("7. PRESCRIÇÕES DIVERSAS", fields.get('prescricoes', '')),
        ("REFERÊNCIAS", fields.get('referencias', ''))
    ]

    for tit, conteudo in secoes:
        p_sec = doc.add_paragraph()
        p_sec.paragraph_format.space_before = Pt(14)
        p_sec.paragraph_format.space_after = Pt(4)
        r_sec = p_sec.add_run(tit)
        r_sec.bold = True
        r_sec.font.name = "Arial"
        r_sec.font.size = Pt(11)

        eh_prescricoes = "PRESCRIÇÕES" in tit.upper()
        renderizar_conteudo_docx(doc, conteudo, forcar_texto=eh_prescricoes)

    p_ass = doc.add_paragraph()
    p_ass.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_ass.paragraph_format.space_before = Pt(28)
    p_ass.paragraph_format.space_after = Pt(4)

    r_ass1 = p_ass.add_run("(assinado eletronicamente)\n")
    r_ass1.italic = True
    r_ass1.font.name = "Arial"
    r_ass1.font.size = Pt(8.5)

    r_ass2 = p_ass.add_run(f"{fields.get('nome_comandante', 'Ten.-Cel. QOEM PM Helder de Lima Dantas Junior')},\n")
    r_ass2.bold = True
    r_ass2.font.name = "Arial"
    r_ass2.font.size = Pt(10.5)

    r_ass3 = p_ass.add_run(f"{fields.get('cargo_comandante', 'Comandante do 18º BPM.')}")
    r_ass3.bold = True
    r_ass3.font.name = "Arial"
    r_ass3.font.size = Pt(10)

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()

def renderizar_conteudo_pdf(story, conteudo, style_subnum, style_body, style_table_hdr, forcar_texto=False):
    conteudo_formatado = formatar_quebras_de_secao(conteudo)
    if not conteudo_formatado or not conteudo_formatado.strip():
        return
    if forcar_texto:
        conteudo_formatado = conteudo_formatado.replace('|', ' ')

    linhas = conteudo_formatado.strip().split('\n')
    i = 0
    while i < len(linhas):
        linha = linhas[i].strip()
        if not linha:
            i += 1
            continue
            
        if not forcar_texto and '|' in linha:
            tabela_linhas = []
            while i < len(linhas) and '|' in linhas[i]:
                cels = [c.strip() for c in linhas[i].split('|')]
                if len(cels) > 1 and cels == "":
                    cels = cels[1:]
                if len(cels) > 1 and cels[-1] == "":
                    cels = cels[:-1]
                if any(c for c in cels):
                    tabela_linhas.append(cels)
                i += 1
                
            if tabela_linhas:
                max_cols = max(len(r) for r in tabela_linhas)
                if max_cols < 2:
                    for r_data in tabela_linhas:
                        txt_line = " ".join([c for c in r_data if c])
                        if txt_line.strip():
                            story.append(Paragraph(txt_line.strip().replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;'), style_body))
                else:
                    col_w = (6.4 * inch) / max_cols
                    
                    pdf_table_data = []
                    for r_idx, r_data in enumerate(tabela_linhas):
                        row_p = []
                        for c_val in r_data:
                            c_clean = c_val.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')
                            if r_idx == 0:
                                p_cell = Paragraph(f"<b><font color='white'>{c_clean}</font></b>", style_table_hdr)
                            else:
                                p_cell = Paragraph(c_clean, style_body)
                            row_p.append(p_cell)
                        pdf_table_data.append(row_p)
                    
                    tbl = Table(pdf_table_data, colWidths=[col_w]*max_cols)
                    tbl.setStyle(TableStyle([
                        ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#002060')),
                        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#A0AAB5')),
                        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
                        ('TOPPADDING', (0,0), (-1,-1), 4),
                        ('BOTTOMPADDING', (0,0), (-1,-1), 4),
                        ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, colors.HexColor('#F8FAFC')])
                    ]))
                    story.append(Spacer(1, 4))
                    story.append(tbl)
                    story.append(Spacer(1, 6))
            continue
            
        l_clean = linha.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')

        if eh_titulo_subsecao(linha):
            story.append(Paragraph(f"<b>{l_clean}</b>", style_subnum))
        else:
            m_item = re.match(r'^([a-z0-9]{1,3}[\.\)\-]|[A-Z][\.\)\-])\s+(.*)', linha, re.IGNORECASE)
            if m_item:
                prefix, rest = m_item.groups()
                p_pre = prefix.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')
                p_rest = rest.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')
                story.append(Paragraph(f"<b>{p_pre}</b> {p_rest}", style_body))
            else:
                story.append(Paragraph(l_clean, style_body))
            
        i += 1

def gerar_ordem_servico_pdf(fields):
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        rightMargin=0.75 * inch,
        leftMargin=0.75 * inch,
        topMargin=0.75 * inch,
        bottomMargin=0.75 * inch
    )

    styles = getSampleStyleSheet()

    style_hdr_left = ParagraphStyle('HdrLeft', parent=styles['Normal'], fontName='Helvetica-Bold', fontSize=10, leading=13, textColor=colors.black)
    style_hdr_right = ParagraphStyle('HdrRight', parent=styles['Normal'], fontName='Helvetica-Bold', fontSize=10, leading=13, alignment=2, textColor=colors.black)
    style_title = ParagraphStyle('OpTitle', parent=styles['Normal'], fontName='Helvetica-Bold', fontSize=12, leading=16, alignment=1, textColor=colors.HexColor('#002060'))
    style_sec_title = ParagraphStyle('SecTitle', parent=styles['Normal'], fontName='Helvetica-Bold', fontSize=11, leading=15, textColor=colors.black, spaceBefore=14, spaceAfter=4)
    style_subnum_title = ParagraphStyle('SubNumTitle', parent=styles['Normal'], fontName='Helvetica-Bold', fontSize=10.5, leading=14, textColor=colors.black, spaceBefore=10, spaceAfter=4)
    style_body = ParagraphStyle('BodyTextCustom', parent=styles['Normal'], fontName='Helvetica', fontSize=10, leading=14, textColor=colors.black, spaceAfter=5)
    style_table_hdr = ParagraphStyle('TableHdrCustom', parent=styles['Normal'], fontName='Helvetica-Bold', fontSize=9.5, leading=13, textColor=colors.white)

    style_ass_small = ParagraphStyle('AssinaturaSmall', parent=styles['Normal'], fontName='Helvetica-Oblique', fontSize=8.5, leading=11, alignment=1, textColor=colors.black)
    style_ass = ParagraphStyle('Assinatura', parent=styles['Normal'], fontName='Helvetica', fontSize=10, leading=14, alignment=1, textColor=colors.black)

    story = []

    left_p = Paragraph("PMPR<br/>2º CRPM/18º BPM<br/>P/3", style_hdr_left)
    right_p = Paragraph(f"Cornélio Procópio, PR.<br/>Em {fields.get('data_expedicao', '')}<br/><b>ORDEM DE SERVIÇO Nº {fields.get('num_os', '077')}</b>", style_hdr_right)

    tbl_hdr = Table([[left_p, right_p]], colWidths=[3.2*inch, 3.2*inch])
    tbl_hdr.setStyle(TableStyle([
        ('VALIGN', (0,0), (-1,-1), 'TOP'),
        ('LEFTPADDING', (0,0), (-1,-1), 0),
        ('RIGHTPADDING', (0,0), (-1,-1), 0),
        ('BOTTOMPADDING', (0,0), (-1,-1), 0),
        ('TOPPADDING', (0,0), (-1,-1), 0),
    ]))

    story.append(tbl_hdr)
    story.append(Spacer(1, 4))
    story.append(HRFlowable(width="100%", thickness=1, color=colors.black, spaceBefore=4, spaceAfter=12))

    op_name = fields.get('nome_operacao', '').strip().upper()
    story.append(Paragraph(f"“{op_name}”", style_title))
    story.append(Spacer(1, 10))

    secoes = [
        ("1. FINALIDADE", fields.get('finalidade', '')),
        ("2. INFORMAÇÕES GERAIS", fields.get('informacoes_gerais', '')),
        ("3. MISSÃO", fields.get('missao', '')),
        ("4. EXECUÇÃO", fields.get('execucao', '')),
        ("5. ADMINISTRAÇÃO E LOGÍSTICA", fields.get('logistica', '')),
        ("6. RELATÓRIOS E SISGCOP", fields.get('relatorios', '')),
        ("7. PRESCRIÇÕES DIVERSAS", fields.get('prescricoes', '')),
        ("REFERÊNCIAS", fields.get('referencias', ''))
    ]

    for tit, conteudo in secoes:
        story.append(Paragraph(tit, style_sec_title))
        eh_prescricoes = "PRESCRIÇÕES" in tit.upper()
        renderizar_conteudo_pdf(story, conteudo, style_subnum_title, style_body, style_table_hdr, forcar_texto=eh_prescricoes)

    story.append(Spacer(1, 20))
    story.append(Paragraph("<i>(assinado eletronicamente)</i>", style_ass_small))
    story.append(Paragraph(f"<b>{fields.get('nome_comandante', 'Ten.-Cel. QOEM PM Helder de Lima Dantas Junior')}</b>,", style_ass))
    story.append(Paragraph(f"<b>{fields.get('cargo_comandante', 'Comandante do 18º BPM.')}</b>", style_ass))

    doc.build(story)
    buffer.seek(0)
    return buffer.getvalue()

def render_ordem_servico():
    col_back, _ = st.columns(2)
    with col_back:
        if st.button("⬅️ Voltar ao Início", key="btn_voltar_os"):
            st.session_state.pagina = "inicio"
            st.rerun()

    if caminho_brasao:
        col1, col2, col3 = st.columns(3)
        with col2:
            st.image(caminho_brasao, width=150)

    st.title("📑 Gerador de Ordem de Serviço (OS)")
    st.caption("18º Batalhão de Polícia Militar — PMPR")
    st.write("Envie a **Ordem de Operação (OO)** para extração automática e geração da **Ordem de Serviço (OS)**.")

    arquivo_oo = st.file_uploader("Envie a Ordem de Operação (PDF ou DOCX)", type=["pdf", "docx", "doc"])
    if arquivo_oo:
        texto_extraido = extrair_texto_arquivo(arquivo_oo)
        if texto_extraido.strip():
            st.success("Ordem de Operação analisada com sucesso!")
            
            parsed_data = parsear_ordem_operacao(texto_extraido)
            
            st.subheader("📝 Informações da Ordem de Serviço")
            st.write("Ajuste os campos essenciais abaixo antes de gerar os documentos:")

            col_a, col_b = st.columns(2)
            with col_a:
                num_os = st.text_input("Número da Ordem de Serviço (OS)", value="077")
                num_oo = st.text_input("Ordem de Operação de Origem", value=parsed_data.get('num_oo', '000/2026'))
            with col_b:
                data_expedicao = st.text_input("Data de Expedição da OS", value=data_atual_extenso())
                nome_operacao = st.text_input("Nome da Operação", value=parsed_data.get('nome_op', 'OPERAÇÃO POLICIAL'))

            col_c, col_d = st.columns(2)
            with col_c:
                nome_comandante = st.text_input("Comandante / Assinatura", value="Ten.-Cel. QOEM PM Helder de Lima Dantas Junior")
            with col_d:
                cargo_comandante = st.text_input("Cargo / Função", value="Comandante do 18º BPM.")

            st.session_state['parsed_full_data'] = parsed_data

            if st.button("🚀 Gerar Ordem de Serviço (Word e PDF)"):
                fields_final = {
                    'num_os': num_os,
                    'num_oo': num_oo,
                    'data_expedicao': data_expedicao,
                    'nome_operacao': nome_operacao,
                    'finalidade': parsed_data.get('finalidade', ''),
                    'informacoes_gerais': parsed_data.get('informacoes_gerais', ''),
                    'missao': parsed_data.get('missao', ''),
                    'execucao': parsed_data.get('execucao', ''),
                    'logistica': parsed_data.get('logistica', ''),
                    'relatorios': parsed_data.get('relatorios', ''),
                    'prescricoes': parsed_data.get('prescricoes', ''),
                    'referencias': parsed_data.get('referencias', ''),
                    'nome_comandante': nome_comandante,
                    'cargo_comandante': cargo_comandante
                }
                
                docx_bytes = gerar_ordem_servico_docx(fields_final)
                st.session_state['generated_docx'] = docx_bytes
                
                pdf_bytes = gerar_ordem_servico_pdf(fields_final)
                st.session_state['generated_pdf'] = pdf_bytes
                
                st.session_state['filename_base'] = f"OS_{num_os.replace('/', '_')}_{nome_operacao.replace(' ', '_')}"

            if 'generated_docx' in st.session_state:
                st.markdown("---")
                st.subheader("📥 Baixar Arquivo Gerado")
                
                col_d1, col_d2 = st.columns(2)
                with col_d1:
                    st.download_button(
                        label="📄 Baixar em Word (.docx)",
                        data=st.session_state['generated_docx'],
                        file_name=f"{st.session_state['filename_base']}.docx",
                        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
                    )
                
                with col_d2:
                    st.download_button(
                        label="📕 Baixar em PDF (.pdf)",
                        data=st.session_state['generated_pdf'],
                        file_name=f"{st.session_state['filename_base']}.pdf",
                        mime="application/pdf"
                    )
        else:
            st.error("Não foi possível extrair texto do arquivo enviado. Verifique se o PDF ou DOCX contém texto pesquisável.")


# =========================================================
# TELA INICIAL / ROTEAMENTO DE PÁGINAS
# =========================================================
if st.session_state.pagina == "inicio":
    if caminho_brasao:
        col1, col2, col3 = st.columns(3)
        with col2:
            st.image(caminho_brasao, width=150)

    st.markdown("""
        <div style="text-align: center; margin-top: 10px; margin-bottom: 25px;">
            <h1 style="font-size: 2.2rem; margin-bottom: 8px;">🛡️ GERADOR DE RELATÓRIOS 🛡️</h1>
            <h3 style="margin-top: 0px; font-weight: 600; color: #3B82F6 !important;">SEÇÃO PM/3</h3>
            <h4 style="margin-top: -5px; font-weight: 500;">18° BPM</h4>
            <h4 style="margin-top: -5px; font-weight: 500;">2° CRPM</h4>
        </div>
    """, unsafe_allow_html=True)

    st.markdown("---")
    st.write("### 📌 Selecione o sistema desejado:")

    col_b1, col_b2 = st.columns(2)
    with col_b1:
        if st.button("🛡️ Relatório Extrajornada", key="btn_ir_extrajornada"):
            st.session_state.pagina = "extrajornada"
            st.rerun()

    with col_b2:
        if st.button("📑 Ordem de Serviço (OS)", key="btn_ir_os"):
            st.session_state.pagina = "ordem_servico"
            st.rerun()

elif st.session_state.pagina == "extrajornada":
    render_extrajornada()

elif st.session_state.pagina == "ordem_servico":
    render_ordem_servico()
